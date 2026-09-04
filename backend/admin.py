"""Admin: reporting, Xero sync queue (mocked), audit logs, user/role management, account."""
import io
import csv
import random
from fastapi import APIRouter, HTTPException, Depends, Response
from pydantic import BaseModel
from typing import Optional, List
from bson import ObjectId

from core import db, now_utc, clean, cleans
from auth import get_current_user, require_admin, ADMIN_ROLES, hash_password
from emails import TEMPLATE_DEFAULTS, get_template, _apply, _shell, _assert_safe_email

admin_router = APIRouter(prefix="/api")


# ---------------- Reporting ----------------
def _range_on(field, date_from, date_to):
    from datetime import datetime as _dt, timezone as _tz
    q = {}
    if date_from:
        q["$gte"] = _dt.fromisoformat(date_from).replace(tzinfo=_tz.utc)
    if date_to:
        q["$lte"] = _dt.fromisoformat(date_to).replace(hour=23, minute=59, second=59, tzinfo=_tz.utc)
    return {field: q} if q else {}


async def _summary_metrics(date_from, date_to):
    rng = _range_on("created_at", date_from, date_to)
    order_q = {"payment_status": "paid", **rng}
    orders = await db.orders.find(order_q).to_list(5000)
    total_ex = sum(o["totals"]["subtotal_ex_vat"] for o in orders)
    total_vat = sum(o["totals"]["vat_total"] for o in orders)
    total_sales = sum(o["totals"]["total_payable"] for o in orders)
    zero_rated = sum(o["totals"]["subtotal_ex_vat"] for o in orders
                     if any(l.get("vat_relief_applied") for l in o.get("items", [])))
    donations = await db.donations.find({"payment_status": "paid", **rng}).to_list(5000)
    donation_total = sum(d["amount"] for d in donations)
    order_donations = sum(o.get("donation_amount", 0) for o in orders)
    refunds = sum(o.get("refund_amount", 0) for o in orders if o.get("refund_amount"))
    products = await db.products.find().to_list(5000)
    carbon_map = {str(p["_id"]): (p.get("carbon_saving_kg", 0) or 0) for p in products}
    items_saved, carbon = 0, 0.0
    for o in orders:
        for it in o.get("items", []):
            q = it.get("quantity", 1) or 1
            items_saved += q
            carbon += carbon_map.get(it.get("product_id"), 0) * q
    low_stock = [clean(p) for p in products
                 if (p.get("quantity_available", 0) - p.get("quantity_reserved", 0)) <= 1
                 and p.get("status") != "sold"]
    aov = round(total_sales / len(orders), 2) if orders else 0
    bookings = await db.event_bookings.count_documents({"status": "confirmed", **rng})
    downloads = await db.resource_downloads.count_documents(_range_on("at", date_from, date_to))
    subs = await db.subscribers.count_documents({**rng})
    eq_donations = await db.equipment_donations.count_documents({**rng})
    return {
        "total_sales_inc_vat": round(total_sales, 2),
        "total_sales_ex_vat": round(total_ex, 2),
        "total_vat": round(total_vat, 2),
        "zero_rated_sales": round(zero_rated, 2),
        "standard_rated_sales": round(total_ex - zero_rated, 2),
        "donations_total": round(donation_total + order_donations, 2),
        "standalone_donations": round(donation_total, 2),
        "order_donations": round(order_donations, 2),
        "refunds_total": round(refunds, 2),
        "average_order_value": aov,
        "orders_count": len(orders),
        "products_sold": items_saved,
        "low_stock_count": len(low_stock),
        "low_stock_products": low_stock[:20],
        "equipment_saved": items_saved,
        "estimated_carbon_saving_kg": round(carbon, 1),
        "equipment_donation_enquiries": eq_donations,
        "event_bookings": bookings,
        "resource_downloads": downloads,
        "email_signups": subs,
    }


COMPARE_KEYS = ["total_sales_inc_vat", "total_sales_ex_vat", "total_vat", "zero_rated_sales",
                "donations_total", "refunds_total", "average_order_value", "orders_count",
                "equipment_saved", "event_bookings", "resource_downloads", "email_signups"]


# ---------------- Editable email templates ----------------
class EmailTemplateBody(BaseModel):
    subject: str
    body: str


@admin_router.get("/admin/email-templates")
async def list_email_templates(user=Depends(require_admin("content_admin", "super_admin"))):
    overrides = {d["key"]: d for d in await db.email_templates.find().to_list(50)}
    out = []
    for key, base in TEMPLATE_DEFAULTS.items():
        o = overrides.get(key)
        out.append({"key": key, "label": base["label"], "description": base["description"],
                    "variables": base["variables"],
                    "subject": (o or {}).get("subject") or base["subject"],
                    "body": (o or {}).get("body") or base["body"],
                    "default_subject": base["subject"], "default_body": base["body"],
                    "customised": bool(o)})
    return out


@admin_router.put("/admin/email-templates/{key}")
async def save_email_template(key: str, body: EmailTemplateBody, user=Depends(require_admin("content_admin", "super_admin"))):
    if key not in TEMPLATE_DEFAULTS:
        raise HTTPException(404, "Unknown template")
    ctx = TEMPLATE_DEFAULTS[key]["sample"]
    try:
        _assert_safe_email(_apply(body.subject, ctx), _shell(_apply(body.body, ctx)))
    except ValueError as e:
        raise HTTPException(400, f"Template failed email safety checks: {e}")
    await db.email_templates.update_one({"key": key},
        {"$set": {"key": key, "subject": body.subject, "body": body.body, "updated_at": now_utc(), "updated_by": user["email"]}}, upsert=True)
    return {"ok": True}


@admin_router.post("/admin/email-templates/{key}/reset")
async def reset_email_template(key: str, user=Depends(require_admin("content_admin", "super_admin"))):
    await db.email_templates.delete_one({"key": key})
    return {"ok": True, **TEMPLATE_DEFAULTS.get(key, {})}


@admin_router.post("/admin/email-templates/{key}/preview")
async def preview_email_template(key: str, body: EmailTemplateBody, user=Depends(require_admin("content_admin", "super_admin"))):
    if key not in TEMPLATE_DEFAULTS:
        raise HTTPException(404, "Unknown template")
    ctx = TEMPLATE_DEFAULTS[key]["sample"]
    return {"subject": _apply(body.subject, ctx), "html": _shell(_apply(body.body, ctx))}


async def _donations_in_range(date_from, date_to):
    from datetime import datetime as _dt, timezone as _tz
    q = {}
    r = {}
    if date_from:
        r["$gte"] = _dt.fromisoformat(date_from).replace(tzinfo=_tz.utc)
    if date_to:
        r["$lte"] = _dt.fromisoformat(date_to).replace(hour=23, minute=59, second=59, tzinfo=_tz.utc)
    if r:
        q["created_at"] = r
    return await db.donations.find(q).sort("created_at", -1).to_list(5000)


def _don_rows(docs):
    def dts(v):
        return v.strftime("%d/%m/%Y") if hasattr(v, "strftime") else str(v or "")
    return [[dts(d.get("created_at")), d.get("reference", ""), d.get("name", ""), d.get("email", ""),
             round(d.get("amount", 0) or 0, 2), "Monthly" if d.get("recurring") else "One-off",
             d.get("payment_status", "")] for d in docs]


DON_COLS = ["Date", "Reference", "Donor", "Email", "Amount (£)", "Type", "Status"]


@admin_router.get("/admin/donations-export.csv")
async def donations_export_csv(date_from: Optional[str] = None, date_to: Optional[str] = None,
                               user=Depends(require_admin("finance_admin"))):
    docs = await _donations_in_range(date_from, date_to)
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(DON_COLS)
    for row in _don_rows(docs):
        w.writerow(row)
    return Response(content=out.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=grace-cares-donations.csv"})


@admin_router.get("/admin/donations-report.xlsx")
async def donations_report_xlsx(date_from: Optional[str] = None, date_to: Optional[str] = None,
                                user=Depends(require_admin("finance_admin"))):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    docs = await _donations_in_range(date_from, date_to)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Donations"
    ws.append(DON_COLS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="006738")
        cell.alignment = Alignment(horizontal="center")
    total = 0.0
    for row in _don_rows(docs):
        ws.append(row)
        total += row[4]
    ws.append([])
    ws.append(["", "", "", "Total", round(total, 2), "", ""])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for i, wdt in enumerate([14, 16, 24, 28, 12, 12, 14], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=grace-cares-donations.xlsx"})


@admin_router.get("/admin/donations-report.pdf")
async def donations_report_pdf(date_from: Optional[str] = None, date_to: Optional[str] = None,
                               user=Depends(require_admin("finance_admin"))):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    docs = await _donations_in_range(date_from, date_to)
    rows = _don_rows(docs)
    total = round(sum(r[4] for r in rows), 2)
    green = colors.HexColor("#006738")
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], textColor=green, fontSize=18, spaceAfter=2)
    sub = ParagraphStyle("sub", parent=ss["Normal"], fontSize=9, textColor=colors.HexColor("#4A4A4D"))
    period = f"{date_from or 'start'} to {date_to or 'now'}" if (date_from or date_to) else "all time"
    data = [DON_COLS] + [[r[0], r[1], r[2], r[3], f"£{r[4]:,.2f}", r[5], r[6]] for r in rows]
    data.append(["", "", "", "Total", f"£{total:,.2f}", "", ""])
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=16 * mm, bottomMargin=16 * mm,
                            leftMargin=14 * mm, rightMargin=14 * mm, title="Grace Cares donations report")
    t = Table(data, colWidths=[22 * mm, 26 * mm, 34 * mm, 44 * mm, 22 * mm, 20 * mm, 20 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), green), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F3F7F4")]),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D9E4DD")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("TEXTCOLOR", (0, -1), (-1, -1), green),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    doc.build([Paragraph("Grace Cares — Financial donations", h1),
               Paragraph(f"Period: {period} · {len(rows)} donation(s) · Total £{total:,.2f}", sub),
               Spacer(1, 8), t])
    return Response(content=buf.getvalue(), media_type="application/pdf",
                    headers={"Content-Disposition": "inline; filename=grace-cares-donations.pdf"})


@admin_router.get("/admin/reports/summary")
async def report_summary(date_from: Optional[str] = None, date_to: Optional[str] = None,
                         user=Depends(require_admin("finance_admin", "readonly", "shop_admin"))):
    cur = await _summary_metrics(date_from, date_to)
    if date_from and date_to:
        from datetime import date as _date, timedelta as _td
        f = _date.fromisoformat(date_from)
        t = _date.fromisoformat(date_to)
        span = (t - f).days + 1
        pf, pt = (f - _td(days=span)).isoformat(), (f - _td(days=1)).isoformat()
        prev = await _summary_metrics(pf, pt)
        cur["previous"] = {k: prev[k] for k in COMPARE_KEYS}
        cur["previous_period"] = {"date_from": pf, "date_to": pt}
    return cur


@admin_router.get("/admin/reports/details")
async def report_details(metric: str, date_from: Optional[str] = None, date_to: Optional[str] = None,
                         user=Depends(require_admin("finance_admin", "readonly", "shop_admin"))):
    rng = _range_on("created_at", date_from, date_to)

    def dts(v):
        return v.strftime("%d/%m/%Y") if hasattr(v, "strftime") else str(v or "")

    if metric in ("total_sales_inc_vat", "total_sales_ex_vat", "total_vat", "orders",
                  "average_order_value", "orders_count", "zero_rated_sales", "refunds_total"):
        q = {"payment_status": "paid", **rng}
        if metric == "refunds_total":
            q["refund_amount"] = {"$gt": 0}
        orders = await db.orders.find(q).sort("created_at", -1).to_list(3000)
        if metric == "zero_rated_sales":
            orders = [o for o in orders if any(l.get("vat_relief_applied") for l in o.get("items", []))]
            cols = ["Date", "Reference", "Customer", "Relieved ex VAT (£)"]
            rows = [[dts(o.get("created_at")), o.get("reference"), o.get("customer", {}).get("name"),
                     round(sum(l["line_ex_vat"] for l in o["items"] if l.get("vat_relief_applied")), 2)] for o in orders]
            return {"title": "Zero-rated (VAT relief) sales", "columns": cols, "rows": rows}
        if metric == "refunds_total":
            cols = ["Date", "Reference", "Customer", "Refund (£)"]
            rows = [[dts(o.get("created_at")), o.get("reference"), o.get("customer", {}).get("name"),
                     o.get("refund_amount", 0)] for o in orders]
            return {"title": "Refunds", "columns": cols, "rows": rows}
        cols = ["Date", "Reference", "Customer", "Ex VAT (£)", "VAT (£)", "Total (£)"]
        rows = [[dts(o.get("created_at")), o.get("reference"), o.get("customer", {}).get("name"),
                 o["totals"]["subtotal_ex_vat"], o["totals"]["vat_total"], o["totals"]["total_payable"]] for o in orders]
        return {"title": "Paid orders", "columns": cols, "rows": rows}

    if metric == "donations_total":
        dons = await db.donations.find({"payment_status": "paid", **rng}).sort("created_at", -1).to_list(3000)
        odons = await db.orders.find({"payment_status": "paid", "donation_amount": {"$gt": 0}, **rng}).sort("created_at", -1).to_list(3000)
        cols = ["Date", "Source", "Donor", "Amount (£)"]
        rows = [[dts(d.get("created_at")), "One-off / monthly", d.get("name") or d.get("email") or "Anonymous", d.get("amount", 0)] for d in dons]
        rows += [[dts(o.get("created_at")), f"With order {o.get('reference')}", o.get("customer", {}).get("name"), o.get("donation_amount", 0)] for o in odons]
        return {"title": "Donations", "columns": cols, "rows": rows}

    if metric == "items_reused":
        orders = await db.orders.find({"payment_status": "paid", **rng}).sort("created_at", -1).to_list(3000)
        cols = ["Date", "Reference", "Item", "SKU", "Qty"]
        rows = []
        for o in orders:
            for it in o.get("items", []):
                rows.append([dts(o.get("created_at")), o.get("reference"), it.get("name"), it.get("sku"), it.get("quantity", 1)])
        return {"title": "Items reused", "columns": cols, "rows": rows}

    if metric == "event_bookings":
        bs = await db.event_bookings.find({"status": "confirmed", **rng}).sort("created_at", -1).to_list(3000)
        cols = ["Date", "Event", "Name", "Attendees"]
        rows = [[dts(b.get("created_at")), b.get("event_name"), b.get("name"), b.get("num_attendees", 1)] for b in bs]
        return {"title": "Event bookings", "columns": cols, "rows": rows}

    if metric == "resource_downloads":
        ds = await db.resource_downloads.find(_range_on("at", date_from, date_to)).sort("at", -1).to_list(3000)
        rmap = {str(r["_id"]): (r.get("title") or r.get("name") or "") for r in await db.resources.find().to_list(500)}
        cols = ["Date", "Resource", "Email"]
        rows = [[dts(d.get("at")), rmap.get(d.get("resource_id"), d.get("resource_id", "")), d.get("email", "")] for d in ds]
        return {"title": "Resource downloads", "columns": cols, "rows": rows}

    if metric == "email_signups":
        subs = await db.subscribers.find({**rng}).sort("created_at", -1).to_list(3000)
        cols = ["Date", "Email", "Source"]
        rows = [[dts(x.get("created_at")), x.get("email"), x.get("consent_source", "")] for x in subs]
        return {"title": "Email signups", "columns": cols, "rows": rows}

    if metric == "low_stock_count":
        products = await db.products.find().to_list(5000)
        low = [p for p in products if (p.get("quantity_available", 0) - p.get("quantity_reserved", 0)) <= 1 and p.get("status") != "sold"]
        cols = ["SKU", "Name", "Available"]
        rows = [[p.get("sku"), p.get("name"), max(0, (p.get("quantity_available", 0) - p.get("quantity_reserved", 0)))] for p in low]
        return {"title": "Low stock", "columns": cols, "rows": rows}

    raise HTTPException(400, "Unknown metric")


@admin_router.get("/admin/reports/orders.csv")
async def orders_csv(user=Depends(require_admin("finance_admin", "readonly"))):
    orders = await db.orders.find().sort("created_at", -1).to_list(5000)
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["Reference", "Date", "Customer", "Status", "Ex VAT", "VAT", "Delivery",
                "Donation", "Total", "Xero Status"])
    for o in orders:
        t = o.get("totals", {})
        w.writerow([o.get("reference"), str(o.get("created_at")), o.get("customer", {}).get("name"),
                    o.get("status"), t.get("subtotal_ex_vat"), t.get("vat_total"),
                    t.get("delivery_total"), t.get("donation"), t.get("total_payable"),
                    o.get("xero_sync_status")])
    return Response(content=out.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=orders.csv"})


# ---------------- Xero sync (mocked) ----------------
@admin_router.get("/admin/xero/queue")
async def xero_queue(user=Depends(require_admin("finance_admin"))):
    items = await db.xero_sync_queue.find().sort("created_at", -1).to_list(500)
    queued = [i for i in items if i.get("status") == "queued"]
    synced = [i for i in items if i.get("status") == "synced"]
    failed = [i for i in items if i.get("status") == "failed"]
    return {"queued": cleans(queued), "synced": cleans(synced), "failed": cleans(failed),
            "counts": {"queued": len(queued), "synced": len(synced), "failed": len(failed)}}


@admin_router.post("/admin/xero/sync/{item_id}")
async def xero_sync_item(item_id: str, user=Depends(require_admin("finance_admin"))):
    item = await db.xero_sync_queue.find_one({"_id": ObjectId(item_id)})
    if not item:
        raise HTTPException(404, "Sync item not found")
    # MOCKED Xero posting: 85% success, idempotent (won't re-post already synced)
    if item.get("status") == "synced":
        return {"ok": True, "already_synced": True}
    attempts = item.get("attempts", 0) + 1
    success = random.random() < 0.85
    if success:
        await db.xero_sync_queue.update_one({"_id": ObjectId(item_id)},
            {"$set": {"status": "synced", "attempts": attempts, "synced_at": now_utc(),
                      "xero_id": f"XERO-{random.randint(10000,99999)}", "error": None}})
        _id = item.get("order_id") or item.get("donation_id")
        if item.get("order_id"):
            await db.orders.update_one({"_id": ObjectId(item["order_id"])},
                                       {"$set": {"xero_sync_status": "synced"}})
        return {"ok": True, "synced": True}
    else:
        await db.xero_sync_queue.update_one({"_id": ObjectId(item_id)},
            {"$set": {"status": "failed", "attempts": attempts,
                      "error": "Xero API temporarily unavailable (503). Retry queued."}})
        return {"ok": False, "failed": True, "error": "Xero temporarily unavailable"}


@admin_router.post("/admin/xero/sync-all")
async def xero_sync_all(user=Depends(require_admin("finance_admin"))):
    items = await db.xero_sync_queue.find({"status": {"$in": ["queued", "failed"]}}).to_list(500)
    results = {"synced": 0, "failed": 0}
    for item in items:
        r = await xero_sync_item(str(item["_id"]), user)
        if r.get("synced"):
            results["synced"] += 1
        else:
            results["failed"] += 1
    return results


@admin_router.get("/admin/xero/reconciliation")
async def reconciliation(user=Depends(require_admin("finance_admin"))):
    orders = await db.orders.find({"payment_status": "paid"}).to_list(5000)
    website_total = sum(o["totals"]["total_payable"] for o in orders)
    synced = await db.xero_sync_queue.find({"status": "synced"}).to_list(5000)
    xero_total = sum(i.get("amount", 0) for i in synced)
    return {"website_total": round(website_total, 2), "xero_total": round(xero_total, 2),
            "difference": round(website_total - xero_total, 2),
            "stripe_total": round(website_total, 2)}


# ---------------- Audit logs ----------------
@admin_router.get("/admin/audit-logs")
async def audit_logs(user=Depends(require_admin("super_admin"))):
    docs = await db.audit_logs.find().sort("at", -1).to_list(500)
    return cleans(docs)


# ---------------- VAT declarations (restricted) ----------------
def _parse_range(date_from, date_to):
    from datetime import datetime as _dt, timezone as _tz
    q = {}
    if date_from:
        q["$gte"] = _dt.fromisoformat(date_from).replace(tzinfo=_tz.utc)
    if date_to:
        q["$lte"] = _dt.fromisoformat(date_to).replace(hour=23, minute=59, second=59, tzinfo=_tz.utc)
    return {"created_at": q} if q else {}


RELIEF_RATE = 0.20  # standard VAT rate that would otherwise apply to relieved items


def _order_relief(order):
    """VAT that would have been charged on relieved lines = the relief granted."""
    relieved = 0.0
    for it in (order or {}).get("items", []):
        if it.get("vat_relief_applied"):
            relieved += round((it.get("line_ex_vat", 0) or 0) * RELIEF_RATE, 2)
    return round(relieved, 2)


async def _orders_by_ref(refs):
    refs = [r for r in refs if r]
    docs = await db.orders.find({"reference": {"$in": refs}}).to_list(len(refs) or 1)
    return {o["reference"]: o for o in docs}


@admin_router.get("/admin/vat-declarations")
async def vat_declarations(date_from: Optional[str] = None, date_to: Optional[str] = None,
                           user=Depends(require_admin("finance_admin"))):
    query = _parse_range(date_from, date_to)
    docs = await db.vat_declarations.find(query).sort("created_at", -1).to_list(2000)
    omap = await _orders_by_ref([d.get("order_reference") for d in docs])
    out = []
    for d in docs:
        c = clean(d)
        o = omap.get(d.get("order_reference"))
        c["total_vat_relieved"] = _order_relief(o)
        c["order_total_paid"] = (o or {}).get("totals", {}).get("total_payable")
        out.append(c)
    return out


@admin_router.get("/admin/vat-declarations/{did}")
async def vat_declaration_detail(did: str, user=Depends(require_admin("finance_admin"))):
    d = await db.vat_declarations.find_one({"_id": ObjectId(did)})
    if not d:
        raise HTTPException(404, "Declaration not found")
    o = await db.orders.find_one({"reference": d.get("order_reference")})
    c = clean(d)
    c["total_vat_relieved"] = _order_relief(o)
    c["order_total_paid"] = (o or {}).get("totals", {}).get("total_payable")
    c["relieved_items"] = [
        {"name": it.get("name"), "sku": it.get("sku"), "quantity": it.get("quantity"),
         "line_ex_vat": it.get("line_ex_vat"),
         "vat_relieved": round((it.get("line_ex_vat", 0) or 0) * RELIEF_RATE, 2)}
        for it in (o or {}).get("items", []) if it.get("vat_relief_applied")
    ]
    return c


VAT_EXPORT_COLUMNS = ["created_at", "order_reference", "customer_email", "eligible_person_name",
                      "eligible_person_address", "condition_description", "for_personal_domestic_use",
                      "completed_by_name", "relationship", "info_accurate", "signature",
                      "order_total_paid", "total_vat_relieved"]


@admin_router.get("/admin/vat-declarations-export.csv")
async def vat_declarations_export(date_from: Optional[str] = None, date_to: Optional[str] = None,
                                  user=Depends(require_admin("finance_admin"))):
    query = _parse_range(date_from, date_to)
    docs = await db.vat_declarations.find(query).sort("created_at", -1).to_list(5000)
    omap = await _orders_by_ref([d.get("order_reference") for d in docs])
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(VAT_EXPORT_COLUMNS)
    for d in docs:
        ca = d.get("created_at")
        o = omap.get(d.get("order_reference"))
        w.writerow([ca.isoformat() if hasattr(ca, "isoformat") else (ca or ""),
                    d.get("order_reference", ""), d.get("customer_email", ""),
                    d.get("eligible_person_name", ""), d.get("eligible_person_address", ""),
                    d.get("condition_description", ""), d.get("for_personal_domestic_use", ""),
                    d.get("completed_by_name", ""), d.get("relationship", ""),
                    d.get("info_accurate", ""), d.get("signature", ""),
                    (o or {}).get("totals", {}).get("total_payable", ""), _order_relief(o)])
    fn = "vat-declarations"
    if date_from or date_to:
        fn += f"_{date_from or 'start'}_to_{date_to or 'now'}"
    return Response(content=out.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={fn}.csv"})


class ReportViewBody(BaseModel):
    preset: str = "all"
    date_from: Optional[str] = None
    date_to: Optional[str] = None


@admin_router.get("/admin/vat-report-view")
async def get_report_view(user=Depends(require_admin("finance_admin"))):
    s = await db.site_settings.find_one({"key": "vat_report_view"})
    return (s or {}).get("view", {"preset": "all", "date_from": None, "date_to": None})


@admin_router.put("/admin/vat-report-view")
async def set_report_view(body: ReportViewBody, user=Depends(require_admin("finance_admin"))):
    await db.site_settings.update_one({"key": "vat_report_view"},
        {"$set": {"key": "vat_report_view", "view": body.model_dump()}}, upsert=True)
    return {"ok": True, "view": body.model_dump()}


class VatStatementBody(BaseModel):
    statement: dict


@admin_router.get("/admin/vat-declaration-statement")
async def get_vat_statement_admin(user=Depends(require_admin("finance_admin", "content_admin"))):
    from shop import get_vat_statement
    return await get_vat_statement()


@admin_router.put("/admin/vat-declaration-statement")
async def set_vat_statement_admin(body: VatStatementBody, user=Depends(require_admin("finance_admin", "content_admin"))):
    from shop import log_audit, DEFAULT_VAT_STATEMENT
    stmt = {k: body.statement.get(k, DEFAULT_VAT_STATEMENT[k]) for k in DEFAULT_VAT_STATEMENT}
    await db.site_settings.update_one({"key": "vat_declaration"},
        {"$set": {"key": "vat_declaration", "statement": stmt}}, upsert=True)
    await log_audit(user, "update", "vat_declaration_statement", after=stmt)
    return {"ok": True, "statement": stmt}


# ---------------- User / role management ----------------
class RoleBody(BaseModel):
    role: str


@admin_router.get("/admin/users")
async def list_users(user=Depends(require_admin("super_admin"))):
    docs = await db.users.find().sort("created_at", -1).to_list(1000)
    out = []
    for d in docs:
        c = clean(d); c.pop("password_hash", None)
        out.append(c)
    return out


@admin_router.put("/admin/users/{uid}/role")
async def set_role(uid: str, body: RoleBody, user=Depends(require_admin("super_admin"))):
    if body.role != "customer" and body.role not in ADMIN_ROLES:
        raise HTTPException(400, "Invalid role")
    before = await db.users.find_one({"_id": ObjectId(uid)})
    await db.users.update_one({"_id": ObjectId(uid)}, {"$set": {"role": body.role}})
    await db.audit_logs.insert_one({"user_email": user["email"], "action": "set_role",
        "entity": "user", "entity_id": uid, "before": {"role": before.get("role")},
        "after": {"role": body.role}, "at": now_utc()})
    return {"ok": True}


# ---------------- Customer account ----------------
class AddressBody(BaseModel):
    label: str
    line1: str
    line2: Optional[str] = ""
    city: str
    postcode: str


@admin_router.get("/account/addresses")
async def get_addresses(user=Depends(get_current_user)):
    u = await db.users.find_one({"_id": ObjectId(user["id"])})
    return u.get("addresses", [])


@admin_router.post("/account/addresses")
async def add_address(body: AddressBody, user=Depends(get_current_user)):
    await db.users.update_one({"_id": ObjectId(user["id"])},
                              {"$push": {"addresses": body.model_dump()}})
    u = await db.users.find_one({"_id": ObjectId(user["id"])})
    return u.get("addresses", [])


class PrefsBody(BaseModel):
    marketing_consent: bool


@admin_router.put("/account/preferences")
async def update_prefs(body: PrefsBody, user=Depends(get_current_user)):
    await db.users.update_one({"_id": ObjectId(user["id"])},
                              {"$set": {"marketing_consent": body.marketing_consent}})
    return {"ok": True}


@admin_router.get("/account/bookings")
async def my_bookings(user=Depends(get_current_user)):
    docs = await db.event_bookings.find({"email": user["email"]}).sort("created_at", -1).to_list(200)
    return cleans(docs)


@admin_router.post("/account/request-deletion")
async def request_deletion(user=Depends(get_current_user)):
    await db.deletion_requests.insert_one({"user_id": user["id"], "email": user["email"],
                                           "status": "requested", "at": now_utc()})
    return {"ok": True, "message": "Your account deletion request has been received."}
